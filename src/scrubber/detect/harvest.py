"""Pre-pass that collects names the data itself labels as names: email display
names and spreadsheet name/company columns. They become an exact-match list
for free text, which catches names the NER model misses and lets "Maria" map
to the same token as "Maria Gonzalez" when that is unambiguous."""

from __future__ import annotations

import csv
import email.utils
import io
import json
import re
import unicodedata
import zipfile
from email import policy
from email.parser import BytesHeaderParser
from pathlib import Path

from spacy.lang.en.stop_words import STOP_WORDS

from ..formats.columns import classify_header
from ..formats.email_ import ADDRESS_HEADERS, _looks_like_org
from ..formats.mbox import iter_messages
from ..transform.pseudonym import HONORIFICS, SUFFIXES, normalize, normalize_person
from .patterns import COMMON_WORD_NAMES

MAX_ROWS = 2_000_000


def _person_rows(surface: str, source: str) -> list[tuple[str, str, str, str]]:
    surface = re.sub(r"\s+", " ", surface.strip().strip('"\'')).strip()
    if not surface or "@" in surface or len(surface) < 3 or len(surface) > 80 or _looks_like_org(surface):
        return []
    if re.search(r"\d", surface):
        return []
    norm = normalize_person(surface)
    if not norm:
        return []
    rows = [(norm, "PERSON", surface, source)]
    if "," in surface:
        last, _, first = surface.partition(",")
        if first.strip():
            rows.append((norm, "PERSON", f"{first.strip()} {last.strip()}", source))
    return rows


def _org_rows(surface: str, source: str) -> list[tuple[str, str, str, str]]:
    surface = re.sub(r"\s+", " ", str(surface).strip())
    if len(surface.split()) < 2 or len(surface) > 100:
        return []
    return [(normalize("ORGANIZATION", surface), "ORGANIZATION", surface, source)]


def _from_headers(raw: bytes, source: str, parser: BytesHeaderParser) -> list:
    rows = []
    msg = parser.parsebytes(raw)
    for name in ADDRESS_HEADERS:
        for value in msg.get_all(name) or []:
            for display, _addr in email.utils.getaddresses([str(value)]):
                if display:
                    rows += _person_rows(display, source)
    return rows


def harvest_file(kind: str, path: Path) -> list[tuple[str, str, str, str]]:
    source = kind
    rows: list = []
    parser = BytesHeaderParser(policy=policy.default)
    if kind == "eml":
        rows += _from_headers(path.read_bytes(), source, parser)
    elif kind == "mbox":
        with open(path, "rb") as f:
            for raw in iter_messages(f, 0, path.stat().st_size + 1):
                try:
                    rows += _from_headers(raw, source, parser)
                except Exception:
                    continue
    elif kind == "csv":
        with open(path, "rb") as fb:
            sample = fb.read(1 << 16)
        enc = "utf-8-sig" if sample.startswith(b"\xef\xbb\xbf") else "utf-8"
        try:
            sample.decode(enc)
        except UnicodeDecodeError:
            enc = "cp1252"
        with open(path, encoding=enc, errors="replace", newline="") as f:
            delim = "\t" if path.suffix.lower() == ".tsv" else ","
            try:
                delim = csv.Sniffer().sniff(sample.decode(enc, errors="replace"), delimiters=",\t;|").delimiter
            except csv.Error:
                pass
            reader = csv.reader(f, delimiter=delim)
            headers = next(reader, [])
            rows += _harvest_table(headers, reader, source)
    elif kind == "json":
        if path.stat().st_size <= 256 << 20:
            rows += _harvest_json(json.loads(path.read_bytes().decode("utf-8-sig", errors="replace")), source)
    elif kind == "zip":
        with zipfile.ZipFile(path) as z:
            for info in z.infolist():
                name = info.filename.rsplit("/", 1)[-1]
                if name in ("users.json", "org_users.json") and info.file_size <= 256 << 20:
                    rows += _harvest_json(json.loads(z.read(info).decode("utf-8-sig", errors="replace")), "slack")
    elif kind == "xml" and path.stat().st_size <= 256 << 20:
        from xml.etree import ElementTree as ET
        for el in ET.parse(path).getroot().iter():
            if isinstance(el.tag, str) and el.tag.rsplit("}", 1)[-1] == "name":
                given = [c.text for c in el if isinstance(c.tag, str) and c.tag.endswith("given") and c.text]
                family = [c.text for c in el if isinstance(c.tag, str) and c.tag.endswith("family") and c.text]
                if family:
                    rows += _fhir_name_rows({"given": given, "family": family}, "xml")
    elif kind in ("xlsx", "ods"):
        from openpyxl import load_workbook
        if kind == "ods":
            from ..formats import ods
            wb = ods.to_workbook(path)
        else:
            from ..formats.tabular import slim_xlsx
            wb = load_workbook(io.BytesIO(slim_xlsx(path.read_bytes())[0]), read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                it = ws.iter_rows(values_only=True)
                headers = None
                for _ in range(20):
                    row = next(it, None)
                    if row is None:
                        break
                    if sum(1 for v in row if isinstance(v, str) and v.strip()) >= 2:
                        headers = list(row)
                        break
                if headers:
                    rows += _harvest_table(headers, it, source)
        finally:
            wb.close()
    return rows


def _fhir_name_rows(hn: dict, source: str) -> list:
    """A FHIR HumanName. Synthea-style names carry digits ("Abram53 Hirthe744"), so they are
    taken as they are; each part is also listed, since file names and notes use parts alone."""
    given = [g.strip() for g in hn.get("given", []) if isinstance(g, str) and g.strip()]
    fam = hn["family"]
    family = " ".join(x for x in fam if isinstance(x, str)).strip() if isinstance(fam, list) else fam.strip()
    full = " ".join(given + [family])
    if not family or len(full) > 80:
        return []
    norm = normalize_person(full) or full.lower()
    rows = [(norm, "PERSON", full, source)]
    rows += [(norm, "PERSON", part, source) for part in given + [family]
             if len(part) > 2 and (re.search(r"\d", part) or part[:1].isupper())]
    return rows


def _harvest_json(obj, source: str) -> list:
    """Values under person-name keys anywhere in a JSON document (Slack users.json profiles...)."""
    from ..formats.json_ import PERSON_KEYS
    out: list = []
    stack = [obj]
    while stack:
        node = stack.pop()
        if isinstance(node, dict):
            if isinstance(node.get("family"), (str, list)) and isinstance(node.get("given"), list):
                out += _fhir_name_rows(node, source)          # FHIR HumanName (EHR exports)
            first, last = node.get("first_name"), node.get("last_name")
            if isinstance(first, str) and isinstance(last, str) and first.strip() and last.strip():
                out += _person_rows(f"{first.strip()} {last.strip()}", source)
            for k, v in node.items():
                if isinstance(v, str) and k.lower() in PERSON_KEYS and k.lower() not in ("first_name", "last_name"):
                    out += _person_rows(v, source)
                elif isinstance(v, (dict, list)):
                    stack.append(v)
        elif isinstance(node, list):
            stack.extend(node)
    return out


def _harvest_table(headers, rows_iter, source) -> list:
    types = [classify_header(h) for h in headers]
    name_cols = [i for i, t in enumerate(types) if t == "person_name"]
    org_cols = [i for i, t in enumerate(types) if t == "organization"]
    # First/Last split across two columns -> also harvest the combined name
    lowered = [str(h or "").lower() for h in headers]
    first_i = next((i for i, h in enumerate(lowered) if re.search(r"first|given|^fname", h)), None)
    last_i = next((i for i, h in enumerate(lowered) if re.search(r"last|family|surname|^lname", h)), None)
    out = []
    if not name_cols and not org_cols:
        return out
    for n, row in enumerate(rows_iter):
        if n >= MAX_ROWS:
            break
        for i in name_cols:
            if i < len(row) and isinstance(row[i], str):
                out += _person_rows(row[i], source)
        if first_i is not None and last_i is not None and max(first_i, last_i) < len(row):
            first, last = row[first_i], row[last_i]
            if isinstance(first, str) and isinstance(last, str) and first.strip() and last.strip():
                out += _person_rows(f"{first.strip()} {last.strip()}", source)
        for i in org_cols:
            if i < len(row) and isinstance(row[i], str):
                out += _org_rows(row[i], source)
    return out


def known_entries(rows, allow_orgs: set[str]) -> tuple[list[tuple[str, str, bool]], list[str]]:
    """From harvested rows -> (detector match entries, full person names for aliasing)."""
    entries: set[tuple[str, str, bool]] = set()
    full_names: set[str] = set()
    for row in rows:
        norm, etype, surface = row["norm"], row["etype"], row["surface"]
        if etype == "PERSON":
            words = [w for w in re.findall(r"[A-Za-z][A-Za-z'-]+", surface)
                     if w.lower().strip(".") not in HONORIFICS and w.lower().strip(".") not in SUFFIXES]
            if len(words) >= 2:
                entries.add((" ".join(words), "PERSON", True))
                full_names.add(norm)
            for w in words:
                lw = w.lower()
                if len(w) >= 3 and lw not in STOP_WORDS and lw not in COMMON_WORD_NAMES:
                    entries.add((w, "PERSON", True))
            # names written with digits (EHR test data such as Synthea: "Abram53 Hirthe744") match as written
            coded = [unicodedata.normalize("NFC", t) for t in surface.split() if re.fullmatch(r"[^\W\d_][\w'-]*?\d+", t)]
            if coded:
                entries.update((t, "PERSON", True) for t in coded)
                if len(surface.split()) >= 2:
                    entries.add((" ".join(surface.split()), "PERSON", True))
        elif etype == "ORGANIZATION":
            if surface.lower() not in allow_orgs:
                entries.add((surface, "ORGANIZATION", True))
    return sorted(entries), sorted(full_names)
