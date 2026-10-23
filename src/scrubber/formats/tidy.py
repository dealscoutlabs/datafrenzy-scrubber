"""Make a scrubbed patient table read like a hand-made de-identified sheet.

Runs on a worksheet after its cells were scrubbed:
  - people and organizations get plain labels: "Patient 1", "Therapist 3", "Facility A"
    (numbered in this sheet, in row order after shuffling)
  - with a patient column, other ID columns (admission/encounter IDs, MRNs) are removed and
    the patient column becomes "Patient ID" in their place
  - date columns become "... Year"; an admission/discharge pair also gets "Length of Stay (Days)"
  - patient rows are shuffled, so the alphabetical order of the original can't be matched back
  - a totals row at the bottom stays at the bottom and is labelled "Totals"
Sheets with merged cells or formulas inside the data rows are left in their scrubbed layout.
"""

from __future__ import annotations

import datetime as dt
import random
import re
import string
from copy import copy

from .columns import normalize_header

PATIENT_HEADER = re.compile(r"\b(patient|client|member|resident|participant|subscriber|insured)\b|^(full )?name$")
ADMIT = re.compile(r"\b(admit|admitted|admission|intake|check ?in|start)\b")
DISCHARGE = re.compile(r"\b(discharge|discharged|check ?out|release|end)\b")
ORG_NOUNS = ("facility", "clinic", "hospital", "pharmacy", "practice", "employer", "school", "payer", "insurer",
             "carrier", "vendor", "program", "campus", "location", "site", "agency")
ID_TYPES = {"other_id", "mrn"}


def _role(header: str, ctype: str) -> str | None:
    h = normalize_header(header)
    if ctype == "person_name":
        if PATIENT_HEADER.search(h):
            return "Patient"
        words = [w for w in h.split() if w not in ("name", "full", "assigned", "primary", "current", "of", "the")]
        return words[-1].title() if words else "Person"
    if ctype == "organization":
        return next((n.title() for n in ORG_NOUNS if n in h), "Organization")
    return None


def _year_header(header: str) -> str:
    h = str(header)
    new = re.sub(r"(?i)\b(date|dt)\b", "Year", h)
    return new if new != h else f"{h} Year"


def _letters(n: int) -> str:
    out = ""
    n += 1
    while n:
        n, r = divmod(n - 1, 26)
        out = string.ascii_uppercase[r] + out
    return out


def tidy_sheet(ws, hdr: int, headers: list, types: list[str], originals: dict[int, list]) -> set[str]:
    """originals: column index -> the original (pre-scrub) values of that column, one per data row.
    Returns the labels it wrote ("Patient 1", "Facility A"), empty if the sheet was left as it is."""
    if hdr is None or ws.merged_cells.ranges:
        return set()
    ncol = len(headers)
    roles = [_role(h, t) if h else None for h, t in zip(headers, types)]
    patient = next((i for i, r in enumerate(roles) if r == "Patient"), None)
    if patient is None:
        return set()  # other tables keep their tokens, which link the same person across files

    # ---- data rows and a trailing totals row
    rows = list(ws.iter_rows(min_row=hdr + 1, max_row=ws.max_row, max_col=ncol))
    body, tail = rows, []
    while body and any(c.data_type == "f" for c in body[-1]):
        tail.insert(0, body.pop())
    while body and all(c.value in (None, "") for c in body[-1]):
        tail.insert(0, body.pop())
    if any(c.data_type == "f" for row in body for c in row):
        return set()
    values = [[c.value for c in row] for row in body]
    orig = {j: (originals.get(j) or [])[: len(values)] for j in range(ncol)}

    # ---- the output columns: ("col", source index, header) or ("los", admit index, discharge index)
    dates = [j for j in range(ncol) if types[j] == "date"]
    admit = next((j for j in dates if ADMIT.search(normalize_header(headers[j]))), None)
    disch = next((j for j in dates if DISCHARGE.search(normalize_header(headers[j]))), None)
    pair = admit is not None and disch is not None and admit != disch
    drop_ids = {j for j in range(ncol) if patient is not None and types[j] in ID_TYPES}
    plan: list[tuple] = []
    for j in range(ncol):
        if j == patient and drop_ids:
            continue  # moves to where the first ID column was
        if j in drop_ids:
            if j == min(drop_ids):
                plan.append(("col", patient, "Patient ID"))
            continue
        if pair and j in (admit, disch):
            if j == min(admit, disch):
                plan += [("col", admit, _year_header(headers[admit])), ("col", disch, _year_header(headers[disch])),
                         ("los", admit, disch)]
            continue
        header = headers[j]
        if types[j] == "date":
            header = _year_header(header)
        elif j == patient:
            header = "Patient ID"
        elif types[j] == "organization":
            header = roles[j]
        plan.append(("col", j, header))

    # ---- shuffle patient rows, then give plain labels in the new order
    order = list(range(len(values)))
    if patient is not None:
        random.SystemRandom().shuffle(order)
    labels: dict[tuple[str, object], str] = {}
    counters: dict[str, int] = {}

    def label(role: str, token) -> str:
        key = (role, token)
        if key not in labels:
            n = counters.get(role, 0)
            counters[role] = n + 1
            labels[key] = f"{role} {_letters(n) if role in ORG_NOUNS_TITLE or role == 'Organization' else n + 1}"
        return labels[key]

    def los(k: int, a: int, d: int):
        a, d = orig[a][k] if k < len(orig[a]) else None, orig[d][k] if k < len(orig[d]) else None
        if isinstance(a, dt.datetime | dt.date) and isinstance(d, dt.datetime | dt.date):
            days = (d - a).days
            return days if days >= 0 else None
        return None

    out_rows = []
    for k in order:
        row = values[k]
        new = []
        for kind, a, b in plan:
            if kind == "los":
                new.append(los(k, a, b))
                continue
            v = row[a]
            if roles[a] and isinstance(v, str) and v.strip():
                v = label(roles[a], v)
            new.append(v)
        out_rows.append(new)

    # ---- write back, carrying each source column's look along
    width = max(ncol, len(plan))
    head_cells = {j: ws.cell(hdr, j + 1) for j in range(ncol)}
    body_cells = {j: (body[0][j] if body else head_cells[j]) for j in range(ncol)}
    widths = {j: ws.column_dimensions[_col(j)].width for j in range(ncol)}
    head_style = {j: copy(head_cells[j]._style) for j in range(ncol)}
    body_style = {j: copy(body_cells[j]._style) for j in range(ncol)}
    body_fmt = {j: body_cells[j].number_format for j in range(ncol)}
    tail_vals = [[(c.value, c.data_type, copy(c._style)) for c in row] for row in tail]

    for c, (kind, a, b) in enumerate(plan):
        src = a if kind == "col" else next((j for j in range(ncol) if types[j] not in ("date",) and
                                           isinstance(values[0][j] if values else None, int)), a)
        cell = ws.cell(hdr, c + 1)
        cell.value = "Length of Stay (Days)" if kind == "los" else plan[c][2]
        cell._style = copy(head_style[a])
        if widths.get(a if kind == "col" else src):
            ws.column_dimensions[_col(c)].width = widths[a if kind == "col" else src]
        for r, new in enumerate(out_rows):
            cell = ws.cell(hdr + 1 + r, c + 1)
            cell.value = new[c]
            cell._style = copy(body_style[src])
            if isinstance(new[c], int):
                cell.number_format = "0" if kind == "col" and types[a] == "date" else body_fmt[src]
    for c in range(len(plan), width):
        for r in range(hdr, hdr + 1 + len(out_rows) + len(tail)):
            ws.cell(r, c + 1).value = None

    # totals row: formulas move with their column
    from openpyxl.formula.translate import Translator
    first_tail = hdr + 1 + len(out_rows)
    for t, row in enumerate(tail_vals):
        r = first_tail + t
        for c in range(width):
            ws.cell(r, c + 1).value = None
        for c, (kind, a, _) in enumerate(plan):
            if kind != "col":
                continue
            value, dtype, style = row[a]
            if dtype == "f":
                value = Translator(str(value), origin=f"{_col(a)}{r}").translate_formula(f"{_col(c)}{r}")
            ws.cell(r, c + 1).value = value
            ws.cell(r, c + 1)._style = style
        if any(d == "f" for _, d, _ in row) and ws.cell(r, 1).value in (None, ""):
            ws.cell(r, 1).value = "Totals"
    return set(labels.values())


ORG_NOUNS_TITLE = {n.title() for n in ORG_NOUNS}


def _col(j: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(j + 1)
