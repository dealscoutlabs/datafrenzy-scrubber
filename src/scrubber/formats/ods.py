"""OpenDocument spreadsheets (.ods: LibreOffice, OpenOffice, Google Sheets downloads).

Read with the standard library and handed to the .xlsx scrubber as a workbook: cell values,
dates and sheet names come across; formulas keep their last calculated value, and comments,
styles and charts are left out. The scrubbed copy is written as .xlsx.
"""

from __future__ import annotations

import datetime as dt
import io
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

T = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
O = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"
X = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
MAX_REPEAT = 10_000  # a repeated non-empty cell or row beyond this is cut off


def _inline(node) -> str:
    out = [node.text or ""]
    for child in node:
        if child.tag == f"{X}s":
            out.append(" " * int(child.get(f"{X}c", "1")))
        elif child.tag == f"{X}tab":
            out.append("\t")
        elif child.tag == f"{X}line-break":
            out.append("\n")
        elif child.tag != f"{O}annotation":
            out.append(_inline(child))
        out.append(child.tail or "")
    return "".join(out)


def _text(cell) -> str:
    """Text of the cell's paragraphs (comments, in office:annotation, are left out)."""
    return "\n".join(_inline(p) for p in cell.findall(f"{X}p"))


def _value(cell):
    kind = cell.get(f"{O}value-type")
    if kind in ("float", "percentage", "currency"):
        v = float(cell.get(f"{O}value", "0"))
        return int(v) if v.is_integer() else v
    if kind == "date":
        raw = cell.get(f"{O}date-value", "")
        try:
            return dt.datetime.fromisoformat(raw) if "T" in raw else dt.datetime.fromisoformat(raw + "T00:00:00")
        except ValueError:
            return _text(cell) or raw
    if kind == "boolean":
        return cell.get(f"{O}boolean-value") == "true"
    text = _text(cell)
    return text if text != "" else None


def _rows(table):
    for el in table.iter():
        if el.tag != f"{T}table-row":
            continue
        cells = []
        for cell in el:
            if cell.tag not in (f"{T}table-cell", f"{T}covered-table-cell"):
                continue
            v = _value(cell)
            n = int(cell.get(f"{T}number-columns-repeated", "1"))
            cells += [v] * (min(n, MAX_REPEAT) if v is not None else n if n < 1024 else 0)
        while cells and cells[-1] is None:
            cells.pop()
        yield cells, int(el.get(f"{T}number-rows-repeated", "1"))


def to_workbook(path: Path):
    from openpyxl import Workbook

    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("content.xml"))
    wb = Workbook()
    wb.remove(wb.active)
    for table in root.iter(f"{T}table"):
        ws = wb.create_sheet((table.get(f"{T}name") or "Sheet")[:31])
        blank = 0
        for cells, repeat in _rows(table):
            if not cells:
                blank += repeat
                continue
            for _ in range(blank):
                ws.append([])
            blank = 0
            for _ in range(min(repeat, MAX_REPEAT)):
                ws.append(cells)
    if not wb.worksheets:
        wb.create_sheet("Sheet")
    return wb


def to_xlsx_bytes(path: Path) -> bytes:
    out = io.BytesIO()
    to_workbook(path).save(out)
    return out.getvalue()
