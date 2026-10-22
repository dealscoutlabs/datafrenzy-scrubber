"""End to end: an input folder -> scrubbed output folder."""

import json
import os
import re

from openpyxl import load_workbook

import scrubber.formats.base as base


def _build_input(paths, synth):
    people = [synth.person() for _ in range(6)]
    a, b, c = people[:3]
    d = paths.input / "Patients" / f"{a.last} {a.first}"
    d.mkdir(parents=True)
    (d / f"visit {a.visit.isoformat()}.eml").write_bytes(bytes(synth.eml(b, a, csv_attachment=True)))
    (d / "notes.txt").write_text(synth.note(a) + synth.note(b))
    crm = paths.input / "CRM"
    crm.mkdir()
    (crm / "contacts.csv").write_text(synth.csv_text(people), newline="")
    (crm / "contacts.xlsx").write_bytes(synth.xlsx_bytes(people))
    mail = paths.input / "Mail"
    mail.mkdir()
    import mailbox
    box = mailbox.mbox(mail / "Inbox")  # Thunderbird-style mbox without an extension
    for p, q in zip(people, people[1:]):
        box.add(synth.eml(p, q, html=False))
    box.close()
    (paths.input / "voicemail.wav").write_bytes(b"RIFF" + bytes(100))
    from fixtures import synth_images as si
    (paths.input / "Patients" / "badge.jpg").write_bytes(si.jpeg_with_gps(si.id_badge(a)))
    from fixtures.synth_docs import SlackWorkspace, docx_bytes, pdf_bytes, scanned_pdf_bytes
    docs = paths.input / "Documents"
    docs.mkdir()
    (docs / f"{c.last} referral.docx").write_bytes(docx_bytes(synth, [c, a, b]))
    (docs / "summary.pdf").write_bytes(pdf_bytes(synth, [b, c, a], csv_text=synth.csv_text([a])))
    (docs / "scan.pdf").write_bytes(si.scanned_pdf([si.text_image(si.note_lines(synth, b))]))
    (docs / "chest.dcm").write_bytes(si.dicom_bytes(synth, c, a))
    (docs / "export.json").write_text(json.dumps([{"patient_name": p.name, "email": p.email, "phone": p.phone,
                                                    "note": f"Seen {p.visit.isoformat()} for {p.first}"} for p in people]))
    (paths.input / "slack-export.zip").write_bytes(SlackWorkspace(synth, people[:5]).zip_bytes(days=2)[0])
    (paths.input / "blob.bin").write_bytes(os.urandom(64))
    (paths.input / ".DS_Store").write_bytes(bytes(16))
    return people


def _secrets(people):
    out = []
    for p in people:
        out += [p.first, p.last, p.email, p.ssn, p.mrn, p.member_id, p.phone, p.street, p.npi]
    return [s for s in out if len(s) >= 4]


def _all_output_bytes(paths):
    for root, _, files in os.walk(paths.output):
        for f in files:
            yield os.path.join(root, f), open(os.path.join(root, f), "rb").read()


def test_end_to_end(workspace, synth, run_pipeline):
    people = _build_input(workspace, synth)
    runner, summary = run_pipeline(workspace, workers=2)

    released = sorted(p.relative_to(workspace.released).as_posix() for p in workspace.released.rglob("*") if p.is_file())
    assert len(released) == 12, (released, summary)
    assert any(re.fullmatch(r"Patients/PERSON_\d{4}/visit 20\d\d\.eml", r) for r in released), released
    assert "CRM/contacts.csv" in released and "CRM/contacts.xlsx" in released and "Mail/Inbox" in released
    assert {"Documents/summary.pdf", "Documents/export.json", "slack-export.zip"} <= set(released)
    assert any(re.fullmatch(r"Documents/PERSON_\d{4} referral\.docx", r) for r in released), released
    assert {"Patients/badge.jpg", "Documents/scan.pdf", "Documents/chest.dcm"} <= set(released)
    assert not any(r.endswith((".wav", ".bin", ".DS_Store")) for r in released)

    # no identifier from the input survives anywhere in /output (data, file names, reports)
    for path, data in _all_output_bytes(workspace):
        for s in _secrets(people):
            assert s.encode() not in data, f"{s!r} in {path}"
            assert s not in path
    # the mapping table and secret never reach /output
    names = {os.path.basename(p) for p, _ in _all_output_bytes(workspace)}
    assert not names & {"secret.key", "state.db", "state.db-wal"}

    s = json.loads((workspace.reports / "run-summary.json").read_text())
    assert [c["number"] for c in s["identifiers"]] == list(range(1, 19))
    by = {c["number"]: c["removed"] for c in s["identifiers"]}
    for n in (1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 13, 14, 16, 17, 18):
        assert by[n] > 0, (n, by)
    reasons = {q["item"]: q["reason"] for q in s["quarantined"]}
    assert any("voiceprint" in r for r in reasons.values())
    assert any("unknown" in r for r in reasons.values())
    assert s["files"].get("done") == 12 and not s["files"].get("review"), s["review"]
    assert s["date_range_years"] and 1900 < s["date_range_years"][0] <= s["date_range_years"][1] <= 2026

    # every exported file opens in its normal reader
    load_workbook(workspace.released / "CRM" / "contacts.xlsx")
    import docx
    import pymupdf
    import zipfile
    docx.Document(next(workspace.released.glob("Documents/*.docx")))
    assert pymupdf.open(workspace.released / "Documents" / "summary.pdf").page_count == 2
    assert zipfile.ZipFile(workspace.released / "slack-export.zip").testzip() is None


def test_rerun_is_idempotent(workspace, synth, run_pipeline):
    _build_input(workspace, synth)
    run_pipeline(workspace)
    before = {p: d for p, d in _all_output_bytes(workspace) if "/scrubbed/" in p}
    msgs = []
    run_pipeline(workspace, messages=msgs)
    after = {p: d for p, d in _all_output_bytes(workspace) if "/scrubbed/" in p}
    assert before == after
    assert not any(m.startswith("Scrubbing") for m in msgs)


def test_config_change_reprocesses(workspace, synth, run_pipeline):
    _build_input(workspace, synth)
    run_pipeline(workspace)
    msgs = []
    run_pipeline(workspace, overrides={"deny_list": {"OTHER_ID": ["Active"]}}, messages=msgs)
    assert any(m.startswith("Scrubbing") for m in msgs)


def test_failed_verification_goes_to_review_not_output(workspace, synth, run_pipeline, monkeypatch):
    _build_input(workspace, synth)
    real = base.Context.residual

    def flag_notes(self, text, where):
        if where.endswith("notes.txt"):
            self.issue(f"{where}: residual PERSON (score 0.90) after scrubbing")
        return real(self, text, where)

    monkeypatch.setattr(base.Context, "residual", flag_notes)
    runner, summary = run_pipeline(workspace)
    monkeypatch.setattr(base.Context, "residual", real)
    assert summary["files"]["review"] == 1
    assert not list(workspace.released.rglob("notes.txt"))
    pending = runner.store.reviews("pending")
    assert len(pending) == 1
    runner.approve(pending[0]["rel"], by="Compliance Officer", note="checked manually")
    assert len(list(workspace.released.rglob("notes.txt"))) == 1
    s = runner.write_summary()
    assert s["review"][0]["status"] == "approved" and s["review"][0]["decided_by"] == "Compliance Officer"


def test_corrupt_file_quarantined(workspace, run_pipeline):
    (workspace.input / "broken.xlsx").write_bytes(b"not a zip at all")
    _, summary = run_pipeline(workspace)
    assert summary["files"]["quarantined"] == 1
    assert not workspace.released.exists() or not list(workspace.released.rglob("*.xlsx"))
    assert "encrypted or damaged" in summary["quarantined"][0]["reason"]


def test_csv_resume_after_interruption(tmp_path, workspace, synth, run_pipeline, monkeypatch):
    import shutil
    import pytest
    import scrubber.formats.tabular as tab

    people = [synth.person() for _ in range(40)]
    rows = synth.csv_text(people * 30)  # 1200 rows -> 3 batches
    (workspace.input / "big.csv").write_text(rows, newline="")
    ref = type(workspace)(tmp_path / "ri", tmp_path / "ro", tmp_path / "rw", tmp_path / "rs")
    shutil.copytree(workspace.input, ref.input)
    run_pipeline(ref)
    expected = (ref.released / "big.csv").read_bytes()

    workspace.state.mkdir(parents=True, exist_ok=True)
    shutil.copy(ref.state / "secret.key", workspace.state / "secret.key")
    real = tab.CellRewriter.rewrite
    calls = {"n": 0}

    def flaky(self, batch):
        calls["n"] += 1
        if calls["n"] == 3:
            raise KeyboardInterrupt
        return real(self, batch)

    monkeypatch.setattr(tab.CellRewriter, "rewrite", flaky)
    with pytest.raises(KeyboardInterrupt):
        run_pipeline(workspace)
    monkeypatch.setattr(tab.CellRewriter, "rewrite", real)
    runner, _ = run_pipeline(workspace)
    assert runner.store.shards("big.csv")[0]["progress"] > 0
    assert (workspace.released / "big.csv").read_bytes() == expected
