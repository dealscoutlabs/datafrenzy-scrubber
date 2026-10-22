"""Every one of the 18 Safe Harbor identifier cases, in every supported format."""

import io
import json
import zipfile
from email.message import EmailMessage
from email import policy

import pytest

from test_identifiers import CASES

TEXTS = [c[1] for c in CASES]
SECRETS = [(c[1], s) for c in CASES for s in c[2]]


def _check(out_text: str, fmt: str):
    missed = [(case, s) for case, s in SECRETS if s in out_text]
    assert not missed, f"{fmt}: {missed[:5]}"


def test_txt(ctx):
    from scrubber.formats.dispatch import scrub_bytes
    out = scrub_bytes("notes.txt", "\n\n".join(TEXTS).encode(), ctx, "notes.txt").decode()
    _check(out, "txt")


def test_eml(ctx):
    from scrubber.formats.email_ import message_text, parse, scrub_message_bytes
    msg = EmailMessage()
    msg["From"], msg["To"], msg["Subject"] = "a@example.com", "b@example.com", "cases"
    msg.set_content("\n\n".join(TEXTS))
    msg.add_alternative("<html><body>" + "".join(f"<p>{t}</p>" for t in TEXTS) + "</body></html>", subtype="html")
    out = scrub_message_bytes(msg.as_bytes(policy=policy.default), ctx, "cases.eml")
    _check(message_text(parse(out)), "eml")


def test_csv(ctx):
    from scrubber.formats.tabular import scrub_csv_bytes
    import csv
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Notes", "Status"])
    for t in TEXTS:
        w.writerow([t, "Open"])
    _check(scrub_csv_bytes(buf.getvalue().encode(), ctx, "cases.csv").decode(), "csv")


def test_xlsx(ctx):
    from openpyxl import Workbook, load_workbook
    from scrubber.formats.tabular import scrub_xlsx_bytes
    wb = Workbook()
    ws = wb.active
    ws.append(["Notes", "Status"])
    for t in TEXTS:
        ws.append([t, "Open"])
    b = io.BytesIO()
    wb.save(b)
    out = load_workbook(io.BytesIO(scrub_xlsx_bytes(b.getvalue(), ctx, "cases.xlsx")))
    _check("\n".join(str(c.value) for row in out.active.iter_rows() for c in row), "xlsx")


def test_json(ctx):
    from scrubber.formats.json_ import scrub_json_bytes
    data = json.dumps([{"comment": t} for t in TEXTS]).encode()
    _check(scrub_json_bytes(data, ctx, "cases.json").decode(), "json")


def test_slack(ctx, synth):
    from fixtures.synth_docs import SlackWorkspace
    from scrubber.formats.archive import scrub_zip_bytes
    people = [synth.person() for _ in range(2)]
    ws = SlackWorkspace(synth, people)
    data, meta = ws.zip_bytes(days=1)
    z = zipfile.ZipFile(io.BytesIO(data))
    files = {n: z.read(n) for n in z.namelist()}
    day = next(n for n in files if n.startswith(meta["channel"] + "/"))
    import datetime as dt
    files[day] = json.dumps([ws.message(people[0], t, dt.date(2022, 3, 14), n=i) for i, t in enumerate(TEXTS)]).encode()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zz:
        for n, c in files.items():
            zz.writestr(n, c)
    out = zipfile.ZipFile(io.BytesIO(scrub_zip_bytes(buf.getvalue(), ctx, "slack.zip")))
    _check("\n".join(out.read(n).decode() for n in out.namelist()), "slack")


def test_docx_runs(ctx, synth):
    import docx
    from fixtures.synth_docs import docx_bytes
    from scrubber.formats.docx_ import all_text, scrub_docx_bytes
    people = [synth.person() for _ in range(3)]
    out = scrub_docx_bytes(docx_bytes(synth, people, cases=TEXTS), ctx, "cases.docx")
    _check("\n".join(all_text(docx.Document(io.BytesIO(out)))), "docx")


def test_pdf(ctx, synth):
    import pymupdf
    from fixtures.synth_docs import pdf_bytes
    from scrubber.formats.pdf import scrub_pdf_bytes
    people = [synth.person() for _ in range(3)]
    out = scrub_pdf_bytes(pdf_bytes(synth, people, cases=TEXTS), ctx, "cases.pdf")
    doc = pymupdf.open(stream=out, filetype="pdf")
    _check("\n".join(p.get_text(sort=True) for p in doc), "pdf")


def test_17_faces_blurred_inside_documents(ctx, synth):
    """Pictures inside PDFs, Word files and emails get the same OCR + face treatment."""
    import numpy as np
    import pymupdf
    from PIL import Image

    from fixtures.synth_images import id_badge, save
    from scrubber.detect.vision import FaceDetector
    from scrubber.formats.email_ import parse, scrub_message_bytes
    from scrubber.formats.pdf import scrub_pdf_bytes

    p = synth.person()
    badge = save(id_badge(p), "PNG")
    fd = FaceDetector(0.6)
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Staff badge on file", fontsize=11)
    page.insert_image(pymupdf.Rect(50, 70, 400, 230), stream=badge)
    out = pymupdf.open(stream=scrub_pdf_bytes(doc.tobytes(), ctx, "badge.pdf"), filetype="pdf")
    xref = out[0].get_images(full=True)[0][0]
    img = Image.open(io.BytesIO(out.extract_image(xref)["image"])).convert("RGB")
    assert fd.detect(np.array(img), 0.2) == []

    msg = EmailMessage()
    msg["From"], msg["To"], msg["Subject"] = "a@example.com", "b@example.com", "badge"
    msg.set_content("attached")
    msg.add_attachment(badge, maintype="image", subtype="png", filename=f"{p.last}_badge.png")
    m = parse(scrub_message_bytes(msg.as_bytes(policy=policy.default), ctx, "badge.eml"))
    att = next(m.iter_attachments())
    assert p.last not in att.get_filename()
    img = Image.open(io.BytesIO(att.get_payload(decode=True))).convert("RGB")
    assert fd.detect(np.array(img), 0.2) == []
