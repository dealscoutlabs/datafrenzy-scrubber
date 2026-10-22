import io
import re
import zipfile

import docx
from docx.shared import Pt

from conftest import assert_no_secrets
from fixtures.synth_docs import docx_bytes
from scrubber.formats.docx_ import scrub_docx_bytes


def _scrub(synth, harvested):
    people = [synth.person() for _ in range(3)]
    src = docx_bytes(synth, people)
    ctx = harvested(people)
    out = scrub_docx_bytes(src, ctx, "visit.docx")
    return people, src, out, ctx


def _xml_text(data: bytes) -> str:
    z = zipfile.ZipFile(io.BytesIO(data))
    return "\n".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith((".xml", ".rels")))


def test_every_part_scrubbed(synth, harvested):
    people, src, out, ctx = _scrub(synth, harvested)
    xml = _xml_text(out)
    # XML escapes & as &amp; - compare on unescaped text
    assert_no_secrets(xml.replace("&amp;", "&"), people, "in docx XML")
    assert ctx.tally.issues == []
    for part in ("footnotes.xml", "comments.xml", "header1.xml", "footer1.xml"):
        text = zipfile.ZipFile(io.BytesIO(out)).read(f"word/{part}").decode()
        assert re.search(r"(PERSON|PHONE|LICENSE|EMAIL)_\d{4}", text), part


def test_opens_and_keeps_formatting(synth, harvested):
    people, src, out, ctx = _scrub(synth, harvested)
    a, b = docx.Document(io.BytesIO(src)), docx.Document(io.BytesIO(out))
    assert len(a.paragraphs) == len(b.paragraphs)
    assert len(a.tables) == len(b.tables) and len(b.tables[0].rows) == 2
    para = next(p for p in b.paragraphs if p.text.startswith("Patient:"))
    runs = para.runs
    assert runs[0].text == "Patient: " and runs[0].bold
    assert re.fullmatch(r"PERSON_\d{4}", runs[1].text) and runs[1].bold and runs[1].font.size == Pt(14)
    assert runs[2].text == "" or not runs[2].text.strip().startswith(people[1].last)
    assert re.fullmatch(r"Patient: PERSON_\d{4}, MRN MRN_\d{4}, DOB (\d{4}|90\+)\.", para.text)
    assert b.tables[0].cell(0, 0).text == "Name"


def test_tracked_changes_comments_and_properties(synth, harvested):
    people, src, out, ctx = _scrub(synth, harvested)
    xml = zipfile.ZipFile(io.BytesIO(out)).read("word/document.xml").decode()
    ins = re.search(r"<w:ins [^>]*w:author=\"([^\"]+)\"[^>]*w:date=\"([^\"]+)\"", xml)
    assert re.fullmatch(r"PERSON_\d{4}", ins.group(1)) and ins.group(2).endswith("-01-01T00:00:00Z")
    deleted = re.search(r"<w:delText[^>]*>([^<]*)</w:delText>", xml).group(1)
    assert "PERSON_" in deleted and "PHONE_" in deleted
    d = docx.Document(io.BytesIO(out))
    cp = d.core_properties
    assert cp.author == "Data Frenzy Scrubber" and cp.last_modified_by == "Data Frenzy Scrubber"
    assert (cp.created.month, cp.created.day) == (1, 1)
    assert "PERSON_" in cp.title
    comments = zipfile.ZipFile(io.BytesIO(out)).read("word/comments.xml").decode()
    assert 'w:initials=""' in comments


def test_hyperlinks_and_images(synth, harvested):
    people, src, out, ctx = _scrub(synth, harvested)
    rels = zipfile.ZipFile(io.BytesIO(out)).read("word/_rels/document.xml.rels").decode()
    assert re.search(r'Target="mailto:email_\d{4}@scrubbed\.invalid"', rels)
    z_in, z_out = zipfile.ZipFile(io.BytesIO(src)), zipfile.ZipFile(io.BytesIO(out))
    media = [n for n in z_in.namelist() if n.startswith("word/media/")]
    assert media and all(z_in.read(n) != z_out.read(n) for n in media)
    from PIL import Image
    with Image.open(io.BytesIO(z_out.read(media[0]))) as im:
        assert im.size == (40, 30) and len(set(im.getdata())) == 1  # blank, same size
