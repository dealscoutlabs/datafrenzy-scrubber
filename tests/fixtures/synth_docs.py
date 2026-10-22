"""Synthetic Slack exports, Word documents and PDFs (Faker identities only)."""

from __future__ import annotations

import datetime as dt
import io
import json
import re
import zipfile

from .synth import Person, Synth


def epoch(d: dt.date, seconds: int = 36000) -> int:
    return int(dt.datetime(d.year, d.month, d.day, tzinfo=dt.timezone.utc).timestamp()) + seconds


def slack_id(s: Synth, prefix: str) -> str:
    return f"{prefix}{s.digits(2)}{s.letters(4)}{s.digits(3)}"


# ------------------------------------------------------------------ Slack

class SlackWorkspace:
    def __init__(self, s: Synth, people: list[Person]):
        self.s = s
        self.people = people
        self.uid = {p.name: slack_id(s, "U") for p in people}
        self.handle = {p.name: f"{p.first.lower()}.{p.last.lower()}" for p in people}
        self.team = slack_id(s, "T")

    def users(self) -> list[dict]:
        out = []
        for p in self.people:
            out.append({
                "id": self.uid[p.name], "team_id": self.team, "name": self.handle[p.name], "deleted": False,
                "real_name": p.name, "tz": "America/Chicago", "tz_label": "Central Standard Time", "tz_offset": -21600,
                "is_bot": False, "updated": epoch(p.visit),
                "profile": {"title": f"Care coordinator at {p.employer}", "phone": p.phone, "real_name": p.name,
                            "real_name_normalized": p.name, "display_name": p.first,
                            "display_name_normalized": p.first, "email": p.email, "first_name": p.first,
                            "last_name": p.last, "image_72": f"https://avatars.slack-edge.com/{self.s.digits(6)}.png",
                            "status_text": f"Visiting {p.city}", "status_emoji": ":car:", "avatar_hash": "g1234abcd"},
            })
        return out

    def message(self, author: Person, text: str, day: dt.date, mention: Person | None = None, n: int = 0,
                thread_ts: str | None = None) -> dict:
        ts = f"{epoch(day, 36000 + n * 61)}.{self.s.digits(6)}"
        uid = self.uid[author.name]
        msg = {
            "type": "message", "user": uid, "ts": ts, "client_msg_id": f"{self.s.digits(8)}-{self.s.digits(4)}",
            "text": (f"<@{self.uid[mention.name]}> " if mention else "") + text,
            "team": self.team,
            "user_profile": {"avatar_hash": "abc123", "image_72": "https://avatars.slack-edge.com/x.png",
                             "first_name": author.first, "real_name": author.name, "display_name": author.first,
                             "team": self.team, "name": self.handle[author.name], "is_restricted": False},
            "blocks": [{"type": "rich_text", "block_id": "Ab1", "elements": [{"type": "rich_text_section", "elements": [
                *([{"type": "user", "user_id": self.uid[mention.name]}] if mention else []),
                {"type": "text", "text": text}]}]}],
            "reactions": [{"name": "thumbsup", "users": [uid], "count": 1}],
        }
        if thread_ts:
            msg["thread_ts"] = thread_ts
            msg["parent_user_id"] = uid
        return msg

    def zip_bytes(self, days: int = 3, root: str = "") -> tuple[bytes, dict]:
        s, ppl = self.s, self.people
        a, b = ppl[0], ppl[1]
        chan_id = slack_id(s, "C")
        chan_name = f"{a.last.lower()}-care-team"
        dm_id = slack_id(s, "D")
        channels = [{"id": chan_id, "name": chan_name, "created": epoch(a.visit), "creator": self.uid[a.name],
                     "is_archived": False, "is_general": False, "members": [self.uid[p.name] for p in ppl],
                     "topic": {"value": f"Coordinating care for {b.name}", "creator": self.uid[a.name], "last_set": 0},
                     "purpose": {"value": f"Patients at {a.street}", "creator": self.uid[a.name], "last_set": 0}}]
        dms = [{"id": dm_id, "created": epoch(a.visit), "members": [self.uid[a.name], self.uid[b.name]]}]
        mpim_name = f"mpdm-{self.handle[a.name]}--{self.handle[b.name]}-1"
        mpims = [{"id": slack_id(s, "G"), "name": mpim_name, "created": epoch(a.visit), "creator": self.uid[a.name],
                  "members": [self.uid[a.name], self.uid[b.name]], "topic": {"value": "", "creator": "", "last_set": 0},
                  "purpose": {"value": f"Group DM {a.first} and {b.first}", "creator": "", "last_set": 0}}]
        files = {}
        files[f"{root}users.json"] = json.dumps(self.users(), indent=4)
        files[f"{root}channels.json"] = json.dumps(channels, indent=4)
        files[f"{root}dms.json"] = json.dumps(dms, indent=4)
        files[f"{root}mpims.json"] = json.dumps(mpims, indent=4)
        files[f"{root}integration_logs.json"] = json.dumps({"logs": []}, indent=4)
        start = dt.date(2022, 3, 14)
        threads = []
        for d in range(days):
            day = start + dt.timedelta(days=d * 9)
            msgs = []
            for n, p in enumerate(ppl[:4]):
                q = ppl[(n + 1) % len(ppl)]
                text = (f"{s.note(q).splitlines()[0]} Her phone is {q.phone}. "
                        f"Details: <mailto:{q.email}|{q.email}> and <{q.url}|chart link>. See <#{chan_id}|{chan_name}>.")
                m = self.message(p, text, day, mention=q, n=n, thread_ts=threads[-1] if threads and n == 3 else None)
                if n == 0:
                    threads.append(m["ts"])
                    m["files"] = [{"id": slack_id(s, "F"), "created": epoch(day), "name": f"{q.last}_labs.pdf",
                                   "title": f"Labs for {q.name}", "mimetype": "application/pdf", "filetype": "pdf",
                                   "user": self.uid[p.name], "url_private": f"https://files.slack.com/{s.digits(8)}/labs.pdf",
                                   "permalink": f"https://acme.slack.com/files/{s.digits(8)}"}]
                    m["edited"] = {"user": self.uid[p.name], "ts": f"{epoch(day, 40000)}.000100"}
                msgs.append(m)
            files[f"{root}{chan_name}/{day.isoformat()}.json"] = json.dumps(msgs, indent=4)
            files[f"{root}{dm_id}/{day.isoformat()}.json"] = json.dumps(
                [self.message(a, f"DOB for {b.first} is {s.date_text(b.dob)}, SSN {b.ssn}", day, n=9)], indent=4)
        files[f"{root}{mpim_name}/{start.isoformat()}.json"] = json.dumps(
            [self.message(b, f"Call {a.name} at {a.phone}", start, n=11)], indent=4)
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for name, content in files.items():
                z.writestr(zipfile.ZipInfo(name, date_time=(2022, 3, 14, 10, 0, 0)), content)
        meta = {"channel": chan_name, "chan_id": chan_id, "dm_id": dm_id, "mpim": mpim_name, "days": days,
                "files": files}
        return buf.getvalue(), meta


# ------------------------------------------------------------------ Word

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def png_bytes(size=(40, 30)) -> bytes:
    from PIL import Image
    out = io.BytesIO()
    Image.new("RGB", size, (200, 30, 30)).save(out, format="PNG")
    return out.getvalue()


def docx_bytes(s: Synth, people: list[Person], cases: list[str] = ()) -> bytes:
    import docx
    from docx.oxml import parse_xml
    from docx.shared import Pt

    a, b, c = people[0], people[1], people[2]
    doc = docx.Document()
    cp = doc.core_properties
    cp.author, cp.last_modified_by = a.name, b.name
    cp.title, cp.subject = f"Visit summary for {b.name}", f"MRN {b.mrn}"
    cp.created = dt.datetime(2022, 5, 17, 9, 30)
    doc.add_heading(f"Visit summary for {b.name}", level=1)
    p = doc.add_paragraph()
    p.add_run("Patient: ").bold = True
    r = p.add_run(b.first)
    r.bold = True
    r.font.size = Pt(14)
    p.add_run(f" {b.last}, MRN {b.mrn}, DOB {s.date_text(b.dob)}.")
    p2 = doc.add_paragraph(f"Seen on {s.date_text(b.visit)} at {b.street}, {b.city}, {b.state} {b.zip}. ")
    p2.add_run(f"Call {b.phone} or fax {b.fax}.").italic = True
    t = doc.add_table(rows=2, cols=3)
    for i, h in enumerate(("Name", "Phone", "SSN")):
        t.cell(0, i).text = h
    t.cell(1, 0).text, t.cell(1, 1).text, t.cell(1, 2).text = c.name, c.phone, c.ssn
    sec = doc.sections[0]
    sec.header.paragraphs[0].text = f"Confidential - prepared by {a.name} ({a.email})"
    sec.footer.paragraphs[0].text = f"Questions: {a.phone}"
    commented = doc.add_paragraph("Insurance verified.")
    doc.add_comment(commented.runs, text=f"Member ID {b.member_id} - spoke with {c.name}", author=c.name,
                    initials="".join(w[0] for w in c.name.split()))
    # tracked changes: an insertion and a deletion carrying names
    tracked = doc.add_paragraph("Emergency contact: ")
    tracked._p.append(parse_xml(
        f'<w:ins xmlns:w="{W}" w:id="91" w:author="{a.name}" w:date="2022-05-18T10:00:00Z">'
        f'<w:r><w:t xml:space="preserve">{b.relative}</w:t></w:r></w:ins>'))
    tracked._p.append(parse_xml(
        f'<w:del xmlns:w="{W}" w:id="92" w:author="{a.name}" w:date="2022-05-18T10:00:00Z">'
        f'<w:r><w:delText xml:space="preserve"> (was {c.name}, {c.phone})</w:delText></w:r></w:del>'))
    link = doc.add_paragraph("Email: ")
    rid = doc.part.relate_to(f"mailto:{b.email}", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                             is_external=True)
    link._p.append(parse_xml(f'<w:hyperlink xmlns:w="{W}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/'
                             f'2006/relationships" r:id="{rid}"><w:r><w:t>{b.email}</w:t></w:r></w:hyperlink>'))
    doc.add_picture(io.BytesIO(png_bytes()))
    for case in cases:
        # split each case into three runs at arbitrary points, so identifiers straddle runs
        cp1, cp2 = max(1, len(case) // 3), max(2, 2 * len(case) // 3)
        para = doc.add_paragraph()
        para.add_run(case[:cp1])
        para.add_run(case[cp1:cp2]).bold = True
        para.add_run(case[cp2:]).italic = True
    out = io.BytesIO()
    doc.save(out)
    return add_footnote(out.getvalue(), f"Footnote: referred by Dr. {c.last}, NPI {c.npi}.")


def add_footnote(data: bytes, text: str) -> bytes:
    """python-docx has no footnote API; add a footnotes part by hand."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    footnotes = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="{W}">'
                 f'<w:footnote w:id="1"><w:p><w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:p></w:footnote>'
                 f'</w:footnotes>')
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            content = src.read(info)
            if info.filename == "[Content_Types].xml":
                content = content.replace(b"</Types>", b'<Override PartName="/word/footnotes.xml" ContentType="'
                                          b'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
            elif info.filename == "word/_rels/document.xml.rels":
                content = content.replace(b"</Relationships>", b'<Relationship Id="rIdFn1" Type="http://schemas.'
                                          b'openxmlformats.org/officeDocument/2006/relationships/footnotes" '
                                          b'Target="footnotes.xml"/></Relationships>')
            elif info.filename == "word/document.xml":
                content = re.sub(rb"(<w:body>)", rb'\1<w:p><w:r><w:t>See note</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r></w:p>',
                                 content, count=1)
            z.writestr(info, content)
        z.writestr("word/footnotes.xml", footnotes)
    return out.getvalue()


# ------------------------------------------------------------------ PDF

def pdf_bytes(s: Synth, people: list[Person], cases: list[str] = (), csv_text: str | None = None) -> bytes:
    import pymupdf

    a, b, c = people[0], people[1], people[2]
    doc = pymupdf.open()
    page = doc.new_page()
    y = 60
    lines = [f"Visit summary for {b.name}", f"MRN: {b.mrn}   DOB: {s.date_text(b.dob)}",
             f"Address: {b.street}, {b.city}, {b.state} {b.zip}", f"Phone {b.phone}, fax {b.fax}, email {b.email}",
             f"SSN {b.ssn}. Insurance member ID: {b.member_id}.", f"Her husband {b.relative} works at {b.employer}.",
             f"Referring provider Dr. {c.last}, NPI {c.npi}."]
    for line in lines + list(cases):
        for sub in line.split("\n"):
            page.insert_text((50, y), sub, fontsize=10)
            y += 15
            if y > 780:
                page = doc.new_page()
                y = 60
    page = doc.new_page()
    page.insert_text((50, 60), f"Page 2: follow up with {a.name} at {a.phone}", fontsize=10)
    first = doc[0]
    annot = first.add_text_annot((400, 60), f"Call {c.name} at {c.phone}")
    annot.set_info(title=a.name, subject=f"Re {b.name}")
    annot.update()
    first.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(50, 50, 300, 62), "uri": b.url})
    w = pymupdf.Widget()
    w.field_name, w.field_type = "patient_name", pymupdf.PDF_WIDGET_TYPE_TEXT
    w.field_value, w.rect = b.name, pymupdf.Rect(50, 700, 250, 720)
    first.add_widget(w)
    first.insert_image(pymupdf.Rect(450, 700, 490, 730), stream=png_bytes())
    if csv_text:
        doc.embfile_add(f"{b.last}_contacts.csv", csv_text.encode(), filename=f"{b.last}_contacts.csv",
                        desc=f"Contacts of {b.name}")
    doc.set_toc([[1, f"Summary for {b.name}", 1], [1, "Follow up", 2]])
    doc.set_metadata({"author": a.name, "title": f"Summary for {b.name}", "subject": b.mrn, "keywords": b.ssn,
                      "creator": "Word", "producer": "Acme PDF", "creationDate": "D:20220517093000"})
    return doc.tobytes()


def scanned_pdf_bytes() -> bytes:
    import pymupdf
    from PIL import Image
    img = io.BytesIO()
    Image.new("L", (850, 1100), 255).save(img, format="PNG")
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(page.rect, stream=img.getvalue())
    return doc.tobytes()
